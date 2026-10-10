import io
import pytest
import pymupdf
import docx

from backend.app.services.extractor import DocumentExtractor
from backend.app.core.errors import (
    UnsupportedFileTypeException,
    EmptyDocumentException,
    DocumentSizeLimitExceededException,
    DocumentExtractionException,
)


@pytest.fixture
def extractor() -> DocumentExtractor:
    return DocumentExtractor(max_file_size_bytes=100_000)


def test_extract_txt_success(extractor: DocumentExtractor):
    """Verify clean TXT extraction with sections and heading detection."""
    content = b"""# Return Policy
Customers may return items within 60 days of purchase in new condition.

# Shipping Guidelines
Standard shipping takes 3 to 5 business days anywhere in the United States.
"""
    result = extractor.extract("policy.txt", content)
    assert result.document_type == "TXT"
    assert result.filename == "policy.txt"
    assert len(result.sections) == 2
    assert "Return Policy" in result.sections[0].section_path
    assert "60 days" in result.sections[0].text
    assert "Shipping Guidelines" in result.sections[1].section_path
    assert "3 to 5 business days" in result.sections[1].text
    assert result.sections[0].page_number is None


def test_extract_txt_empty_raises(extractor: DocumentExtractor):
    """Verify empty text file raises EmptyDocumentException."""
    with pytest.raises(EmptyDocumentException):
        extractor.extract("empty.txt", b"   \n\t  ")


def test_extract_pdf_multi_page_success(extractor: DocumentExtractor):
    """Verify PyMuPDF multi-page extraction preserves 1-indexed page numbers."""
    doc = pymupdf.open()
    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 72), "Title: Return Policy Page 1\nItems must be unworn.")
    # Page 2
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Title: Shipping Policy Page 2\nFree shipping across US.")

    pdf_bytes = doc.tobytes()
    doc.close()

    result = extractor.extract("handbook.pdf", pdf_bytes)
    assert result.document_type == "PDF"
    assert len(result.sections) == 2
    assert result.sections[0].page_number == 1
    assert "Page 1" in result.sections[0].text
    assert result.sections[1].page_number == 2
    assert "Page 2" in result.sections[1].text


def test_extract_pdf_empty_raises(extractor: DocumentExtractor):
    """Verify empty PDF without text raises EmptyDocumentException."""
    doc = pymupdf.open()
    doc.new_page()  # Blank page with no text
    pdf_bytes = doc.tobytes()
    doc.close()

    with pytest.raises(EmptyDocumentException):
        extractor.extract("blank.pdf", pdf_bytes)


def test_extract_pdf_corrupt_raises(extractor: DocumentExtractor):
    """Verify corrupted PDF bytes raise DocumentExtractionException."""
    with pytest.raises(DocumentExtractionException):
        extractor.extract("corrupted.pdf", b"%PDF-1.4 but not valid binary content")


def test_extract_docx_success(extractor: DocumentExtractor):
    """Verify DOCX extraction preserves heading styles and table contents."""
    doc = docx.Document()
    doc.add_heading("Section 1: Warranty", level=1)
    doc.add_paragraph("All footwear has a 90-day manufacturer warranty.")
    doc.add_heading("Section 2: Exceptions", level=1)
    doc.add_paragraph("Normal wear and tear is excluded.")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Status"
    table.rows[1].cells[0].text = "Boots"
    table.rows[1].cells[1].text = "Covered"

    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    result = extractor.extract("warranty.docx", docx_bytes)
    assert result.document_type == "DOCX"
    assert len(result.sections) >= 2
    assert "Warranty" in result.sections[0].section_path
    assert "90-day" in result.sections[0].text
    assert "Exceptions" in result.sections[1].section_path
    assert "Boots | Covered" in result.raw_text
    assert result.sections[0].page_number is None  # DOCX page numbers are not stable


def test_extract_docx_empty_raises(extractor: DocumentExtractor):
    """Verify empty DOCX raises EmptyDocumentException."""
    doc = docx.Document()
    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    with pytest.raises(EmptyDocumentException):
        extractor.extract("empty.docx", docx_bytes)


def test_extract_docx_corrupt_raises(extractor: DocumentExtractor):
    """Verify corrupted DOCX raises DocumentExtractionException."""
    with pytest.raises(DocumentExtractionException):
        extractor.extract("corrupt.docx", b"PK\x03\x04 fake corrupted zip archive")


def test_file_size_exceeded_raises(extractor: DocumentExtractor):
    """Verify file exceeding size limit raises DocumentSizeLimitExceededException."""
    large_payload = b"A" * 150_000
    with pytest.raises(DocumentSizeLimitExceededException):
        extractor.extract("large.txt", large_payload)


def test_unsupported_extension_raises(extractor: DocumentExtractor):
    """Verify unsupported extensions raise UnsupportedFileTypeException."""
    with pytest.raises(UnsupportedFileTypeException):
        extractor.extract("script.py", b"print('hello')")
